from pathlib import Path
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
REPORT = ROOT / "report.docx"
FLOW = ROOT / "assets" / "flowchart.png"
FONT = "/System/Library/Fonts/Supplemental/Arial Unicode.ttf"
BOLD_FONT = "/System/Library/Fonts/Supplemental/Arial Bold.ttf"
if not Path(FONT).exists():
    FONT = "/Library/Fonts/Arial Unicode.ttf"


def font(size, bold=False):
    return ImageFont.truetype(BOLD_FONT if bold else FONT, size)


def make_flowchart():
    image = Image.new("RGB", (1800, 2600), "white")
    draw = ImageDraw.Draw(image)
    black, blue = "#000000", "#0A8FCB"
    regular, strong = font(40), font(40, True)

    def center_text(box, text, fnt=regular):
        words = text.split()
        lines, line = [], ""
        max_w = box[2] - box[0] - 30
        for word in words:
            candidate = (line + " " + word).strip()
            if draw.textbbox((0, 0), candidate, font=fnt)[2] <= max_w:
                line = candidate
            else:
                lines.append(line); line = word
        if line: lines.append(line)
        heights = [draw.textbbox((0, 0), x, font=fnt)[3] for x in lines]
        y = (box[1] + box[3] - sum(heights) - 10 * (len(lines)-1)) / 2
        for line, h in zip(lines, heights):
            x = (box[0] + box[2] - draw.textbbox((0, 0), line, font=fnt)[2]) / 2
            draw.text((x, y), line, fill=black, font=fnt)
            y += h + 10

    def box(cx, y, w, h, text, kind="rect"):
        x1, y1, x2, y2 = cx-w//2, y, cx+w//2, y+h
        if kind == "oval":
            draw.ellipse((x1, y1, x2, y2), outline=blue, width=7)
        elif kind == "diamond":
            draw.polygon([(cx,y1),(x2, y+h//2),(cx,y2),(x1,y+h//2)], outline=blue, fill="white", width=7)
        else:
            draw.rounded_rectangle((x1,y1,x2,y2), radius=18, outline=blue, width=7)
        center_text((x1, y1, x2, y2), text, strong if kind == "oval" else regular)
        return (x1,y1,x2,y2)

    def arrow(x1,y1,x2,y2,label=None):
        draw.line((x1,y1,x2,y2), fill=black, width=5)
        import math
        a=math.atan2(y2-y1,x2-x1)
        for d in (2.55, -2.55):
            draw.line((x2,y2,x2+25*math.cos(a+d),y2+25*math.sin(a+d)), fill=black, width=5)
        if label:
            draw.text(((x1+x2)//2+15,(y1+y2)//2-35), label, fill=black, font=regular)

    c=900
    start=box(c,70,400,100,"Начало","oval")
    get=box(c,230,700,120,"Получить путь к бинарному файлу")
    openf=box(c,420,700,120,"Открыть файл")
    opened=box(c,610,520,190,"Файл открыт?","diamond")
    read=box(c,880,700,120,"Прочитать следующую запись")
    readok=box(c,1070,520,190,"Запись прочитана?","diamond")
    node=box(c,1340,700,120,"Создать элемент односвязного списка")
    tree=box(c,1600,700,120,"Построить BST по высотам")
    mm=box(c,1780,700,120,"Найти минимальную и максимальную высоты")
    stats=box(c,1960,700,120,"Определить число точек и продолжительность")
    out=box(c,2140,700,120,"Вывести и записать statistics.txt")
    free=box(c,2320,700,120,"Освободить список и дерево")
    end=box(c,2470,400,80,"Конец","oval")
    err=box(1500,610,440,120,"Вывести ошибку", "rect")
    arrow(c,start[3],c,get[1]); arrow(c,get[3],c,openf[1]); arrow(c,openf[3],c,opened[1])
    arrow(c,opened[3],c,read[1],"Да")
    arrow(opened[2],705,err[0],670,"Нет"); arrow(err[2],670,1700,2470); arrow(1700,2470,end[2],2470)
    arrow(c,read[3],c,readok[1]); arrow(c,readok[3],c,node[1],"Да")
    # loop back from node to read
    draw.line((node[0],1400,380,1400,380,940,node[0],940), fill=black, width=5)
    arrow(node[0],940,node[0]+1,940)
    # Ветвь «Нет» обходит блок создания узла и переходит к обработке списка.
    arrow(readok[2],1165,1330,1165,"Нет")
    draw.line((1330,1165,1330,1660,tree[2],1660), fill=black, width=5)
    arrow(tree[2],1660,tree[2]-1,1660)
    arrow(c,tree[3],c,mm[1]); arrow(c,mm[3],c,stats[1]); arrow(c,stats[3],c,out[1]); arrow(c,out[3],c,free[1]); arrow(c,free[3],c,end[1])
    image.save(FLOW)


def set_cell_border(cell, color="FFFFFF"):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement('w:tcBorders')
    for edge in ('top','left','bottom','right','insideH','insideV'):
        tag = OxmlElement(f'w:{edge}')
        tag.set(qn('w:val'), 'single'); tag.set(qn('w:sz'), '0'); tag.set(qn('w:color'), color)
        borders.append(tag)
    tcPr.append(borders)


def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr(); shd = OxmlElement('w:shd'); shd.set(qn('w:fill'), fill); tcPr.append(shd)


def set_repeat_table_header(row):
    trPr = row._tr.get_or_add_trPr(); tblHeader = OxmlElement('w:tblHeader'); tblHeader.set(qn('w:val'), 'true'); trPr.append(tblHeader)


def para(doc, text="", bold=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY, before=0, after=6, first=1.25, size=14):
    p = doc.add_paragraph(); p.alignment = align
    pf = p.paragraph_format; pf.first_line_indent = Cm(first) if first else None; pf.space_before = Pt(before); pf.space_after = Pt(after); pf.line_spacing = 1.5; pf.keep_together = True
    r = p.add_run(text); r.bold = bold; r.font.name = "Times New Roman"; r._element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman'); r.font.size = Pt(size)
    return p


def heading(doc, text, level=1):
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER if level==1 else WD_ALIGN_PARAGRAPH.LEFT
    pf=p.paragraph_format; pf.space_before=Pt(12); pf.space_after=Pt(8); pf.keep_with_next=True
    r=p.add_run(text); r.bold=True; r.font.name="Times New Roman"; r._element.rPr.rFonts.set(qn('w:eastAsia'),'Times New Roman'); r.font.size=Pt(14)
    return p


def code_block(doc, path):
    for line in path.read_text(encoding='utf-8').splitlines():
        p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.LEFT
        pf=p.paragraph_format; pf.space_after=Pt(0); pf.line_spacing=1.0
        r=p.add_run(line if line else " "); r.font.name="Courier New"; r._element.rPr.rFonts.set(qn('w:eastAsia'),'Courier New'); r.font.size=Pt(7.5)


def add_page_number(section):
    footer=section.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    run=footer.add_run(); fldChar1=OxmlElement('w:fldChar'); fldChar1.set(qn('w:fldCharType'),'begin'); instrText=OxmlElement('w:instrText'); instrText.set(qn('xml:space'),'preserve'); instrText.text='PAGE'; fldChar2=OxmlElement('w:fldChar'); fldChar2.set(qn('w:fldCharType'),'end'); run._r.append(fldChar1); run._r.append(instrText); run._r.append(fldChar2)


def build_report():
    make_flowchart()
    doc=Document()
    sec=doc.sections[0]; sec.page_width=Cm(21); sec.page_height=Cm(29.7); sec.left_margin=Cm(3); sec.right_margin=Cm(1.5); sec.top_margin=Cm(2); sec.bottom_margin=Cm(2)
    add_page_number(sec)
    styles=doc.styles
    styles['Normal'].font.name='Times New Roman'; styles['Normal']._element.rPr.rFonts.set(qn('w:eastAsia'),'Times New Roman'); styles['Normal'].font.size=Pt(14)

    # Title page
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(5)
    p.add_run().add_picture(str(ROOT/'assets/mai_logo.png'), width=Cm(1.7))
    for text in ["МОСКОВСКИЙ АВИАЦИОННЫЙ ИНСТИТУТ", "(Национальный исследовательский университет)"]:
        p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(2); p.add_run(text).font.name='Times New Roman'
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(52); p.add_run("Кафедра 305").font.name='Times New Roman'
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(22); p.add_run().add_picture(str(ROOT/'assets/department_logo.jpeg'), width=Cm(5.7))
    for text in ["Отчёт по лабораторной работе № 3 на тему:", "«Разработка навигационного анализатора треков»"]:
        p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(16 if 'Отчёт' in text else 0); r=p.add_run(text); r.font.name='Times New Roman'; r.bold='анализатора' in text
    table=doc.add_table(rows=2, cols=2); table.autofit=False; table.columns[0].width=Cm(4); table.columns[1].width=Cm(10)
    table.cell(0,0).text='Выполнил:'; table.cell(1,0).text='Принял:'
    table.cell(0,1).text='Волянский Андрей Петрович\nстудент гр. М3О-243БВ-24'; table.cell(1,1).text='Рычков Александр Сергеевич\nАссистент кафедры 305'
    for row in table.rows:
        for cell in row.cells:
            set_cell_border(cell); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in cell.paragraphs:
                p.alignment=WD_ALIGN_PARAGRAPH.LEFT if cell==row.cells[0] else WD_ALIGN_PARAGRAPH.RIGHT
                for r in p.runs: r.font.name='Times New Roman'; r.font.size=Pt(14)
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(125); p.add_run('Москва, 2026').font.name='Times New Roman'
    doc.add_page_break()

    heading(doc,'ВВЕДЕНИЕ')
    para(doc,'Цель работы — получить практические навыки разработки программы на языке C для обработки полётных данных с применением динамической памяти, односвязного списка и бинарного дерева поиска.')
    para(doc,'В работе требуется прочитать полётные данные из бинарного файла, представить их в виде односвязного списка, построить бинарное дерево по высотам и по свойствам дерева определить минимальную и максимальную высоты. Программа должна вывести количество точек, продолжительность полёта и высоты, а также записать статистику в текстовый файл. Дополнительно предусмотрены выбор входного файла через командную строку и краткое описание Git.')

    heading(doc,'1 РЕАЛИЗАЦИЯ ЗАДАНИЯ')
    heading(doc,'1.1 Формат бинарных данных',2)
    para(doc,'Проверен файл flight_data6.bin. Его размер составляет 256 000 байт. Одна запись имеет размер 16 байт: 4 байта на метку времени и по 4 байта на широту, долготу и высоту. Размер файла кратен 16, поэтому в нём 16 000 полных записей и неполной записи в конце нет.')
    table=doc.add_table(rows=1,cols=3); table.style='Table Grid'; hdr=table.rows[0].cells
    for c,t in zip(hdr,['Поле','Тип','Назначение']): c.text=t; shade(c,'D9EAF7'); set_cell_border(c,'D9D9D9')
    for row in [('timestamp_ms','uint32_t','время от начала маршрута, мс'),('lat_rad','float','широта, радианы'),('lon_rad','float','долгота, радианы'),('alt_m','float','высота, м')]:
        cells=table.add_row().cells
        for c,t in zip(cells,row): c.text=t; set_cell_border(c,'D9D9D9')
    set_repeat_table_header(table.rows[0])
    para(doc,'Первая отметка времени равна 0 мс, последняя — 319 980 мс; шаг между соседними записями стабильно равен 20 мс. Широта находится в диапазоне 45,779329…–45,893919…°, долгота — 28,619242…–28,676539…°, а высота — 95,00849…–109,99453… м. Для наглядного вывода координаты переводятся из радиан в градусы по формуле degrees = radians × 180 / π.')
    heading(doc,'1.2 Односвязный список',2)
    para(doc,'Каждая запись после чтения переносится в структуру FlightData. Узел FlightNode содержит одну такую структуру и указатель next на следующий узел. Структура FlightList хранит head, tail и count. Указатель tail позволяет добавлять очередной элемент в конец за постоянное время, а функция list_free последовательно освобождает все выделенные узлы.')
    heading(doc,'1.3 Бинарное дерево высот',2)
    para(doc,'Для каждой высоты строится узел HeightNode. При вставке меньшая высота направляется в левое поддерево, большая — в правое. Точное повторение высоты не добавляется: для поиска минимума и максимума это не меняет результат. Минимум находится проходом по левым ссылкам, максимум — по правым. Дерево освобождается рекурсивно после вычисления статистики.')
    heading(doc,'1.4 Определение статистики маршрута',2)
    para(doc,'Количество точек берётся из поля count списка. Продолжительность вычисляется как разность последней и первой временных меток: (319 980 − 0) / 1000 = 319,98 с, или 5,33 мин. Значения min и max получаются из левого и правого крайних узлов BST.')
    heading(doc,'1.5 Запись результатов',2)
    para(doc,'Функция write_statistics открывает output/statistics.txt в текстовом режиме и записывает число точек, продолжительность, максимальную и минимальную высоты. Ошибки открытия, чтения, выделения памяти и закрытия файла проверяются.')

    heading(doc,'2 БЛОК СХЕМА ПРОГРАММЫ')
    para(doc,'Основная последовательность обработки, включая проверку открытия файла и цикл чтения записей, показана на рисунке 1.')
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run().add_picture(str(FLOW), width=Cm(13.4))
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run('Рисунок 1 — Блок-схема основной программы').italic=True

    doc.add_page_break()
    heading(doc,'3 РЕЗУЛЬТАТЫ ВЫПОЛНЕНИЯ')
    para(doc,'Программа была собрана командой make и выполнена на приложенном файле data/flight_data6.bin. Сборка завершилась без предупреждений. Ниже приведён реальный вывод финальной программы.')
    out='''Открытие файла: data/flight_data6.bin\nПрочитано записей: 16000\nПервые 10 элементов односвязного списка:\n  [0] time=0 ms, lat=45.836624 deg, lon=28.676538 deg, alt=102.18 m\n  [1] time=20 ms, lat=45.837200 deg, lon=28.676537 deg, alt=100.23 m\n  [2] time=40 ms, lat=45.837769 deg, lon=28.676525 deg, alt=102.43 m\n  [3] time=60 ms, lat=45.838341 deg, lon=28.676508 deg, alt=105.04 m\n  [4] time=80 ms, lat=45.838917 deg, lon=28.676489 deg, alt=101.02 m\n  [5] time=100 ms, lat=45.839485 deg, lon=28.676456 deg, alt=104.45 m\n  [6] time=120 ms, lat=45.840061 deg, lon=28.676424 deg, alt=103.73 m\n  [7] time=140 ms, lat=45.840630 deg, lon=28.676378 deg, alt=105.09 m\n  [8] time=160 ms, lat=45.841206 deg, lon=28.676331 deg, alt=101.68 m\n  [9] time=180 ms, lat=45.841774 deg, lon=28.676277 deg, alt=100.92 m\n\nСтатистика маршрута:\nКоличество точек: 16000\nПродолжительность полёта: 319.98 с (5.33 мин)\nМаксимальная высота: 109.99 м\nМинимальная высота: 95.01 м\nСтатистика сохранена: output/statistics.txt'''
    for line in out.splitlines():
        p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(0); p.paragraph_format.line_spacing=1; r=p.add_run(line if line else ' '); r.font.name='Courier New'; r.font.size=Pt(8)
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run('Рисунок 2 — Консольный вывод программы').italic=True
    para(doc,'Содержимое созданного файла statistics.txt совпадает с показанной статистикой: 16 000 точек, 319,98 с, максимальная высота 109,99 м и минимальная 95,01 м.')

    heading(doc,'4 ДОПОЛНИТЕЛЬНЫЕ ЗАДАНИЯ')
    heading(doc,'4.1 Выбор бинарного файла через терминал',2)
    para(doc,'Путь к бинарному файлу передаётся первым аргументом командной строки: ./build/flight_analyzer data/flight_data6.bin. Если аргумент не передан, программа использует data/flight_data6.bin по умолчанию и выводит соответствующую подсказку. Такой интерфейс не содержит абсолютных путей и работает в macOS и других POSIX-системах.')
    heading(doc,'4.2 Система контроля версий Git',2)
    para(doc,'Git — распределённая система контроля версий. Она фиксирует состояние файлов в коммитах и позволяет видеть историю изменений, возвращаться к прежним версиям, вести параллельную работу в ветках и объединять результаты. В коллективной разработке Git уменьшает риск потери изменений, упрощает проверку кода и обмен результатами через удалённый репозиторий.')
    heading(doc,'4.3 Основные команды Git',2)
    table=doc.add_table(rows=1,cols=2); table.style='Table Grid'
    for c,t in zip(table.rows[0].cells,['Команда','Назначение']): c.text=t; shade(c,'D9EAF7'); set_cell_border(c,'D9D9D9')
    commands=[('git init','создать локальный репозиторий'),('git clone <URL>','получить копию удалённого репозитория'),('git status','показать состояние файлов'),('git add <файл>','подготовить файл к коммиту'),('git commit -m "сообщение"','сохранить подготовленные изменения'),('git branch','показать или создать ветки'),('git switch <ветка> / git checkout <ветка>','переключиться на ветку'),('git merge <ветка>','объединить изменения веток'),('git pull','получить и объединить удалённые изменения'),('git push','отправить локальные коммиты на сервер'),('git log','просмотреть историю коммитов')]
    for a,b in commands:
        cells=table.add_row().cells; cells[0].text=a; cells[1].text=b
        for c in cells: set_cell_border(c,'D9D9D9')
    set_repeat_table_header(table.rows[0])
    heading(doc,'4.4 Ссылка на репозиторий',2)
    para(doc,'[Ссылка на репозиторий]')
    para(doc,'Удалённый репозиторий не создавался без учётной записи пользователя. В README.md приведены команды, которыми после создания репозитория на GitHub можно добавить адрес и отправить локальную историю.')

    heading(doc,'ЗАКЛЮЧЕНИЕ')
    para(doc,'В ходе лабораторной работы разработана программа на C для анализа навигационных данных. Она прочитала 16 000 записей из бинарного файла, сохранила их в односвязном списке, построила BST высот и определила минимальную высоту 95,01 м и максимальную высоту 109,99 м. Рассчитана продолжительность полёта 319,98 с, а статистика сохранена в текстовом файле. Реализованы выбор входного файла через командную строку, Makefile и материалы по Git. Проверка сборки с AddressSanitizer и UndefinedBehaviorSanitizer не выявила ошибок памяти.')

    for label, filename in [('ПРИЛОЖЕНИЕ А — main.c','src/main.c'),('ПРИЛОЖЕНИЕ Б — flight_data.c','src/flight_data.c'),('ПРИЛОЖЕНИЕ В — flight_data.h','include/flight_data.h'),('ПРИЛОЖЕНИЕ Г — Makefile','Makefile')]:
        doc.add_page_break(); heading(doc,label); code_block(doc,ROOT/filename)

    doc.core_properties.title='Разработка навигационного анализатора треков'
    doc.core_properties.author='Волянский Андрей Петрович'
    doc.save(REPORT)


if __name__ == '__main__':
    build_report()
